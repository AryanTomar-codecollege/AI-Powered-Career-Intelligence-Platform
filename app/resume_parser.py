import re
from io import BytesIO
import fitz
from docx import Document
import spacy

# Keep the supported vocabulary small, explicit, and easy to extend.
SKILL_TERMS = {
    "python", "java", "c++", "c", "javascript", "typescript", "sql", "mysql", "postgresql",
    "fastapi", "flask", "django", "react", "html", "css", "git", "docker", "kubernetes",
    "machine learning", "deep learning", "natural language processing", "nlp", "spacy",
    "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "data analysis",
    "data visualization", "power bi", "tableau", "aws", "azure", "gcp", "rest api",
    "linux", "excel", "communication", "leadership", "problem solving"
}
TITLE_TERMS = {
    "software engineer", "software developer", "backend developer", "frontend developer",
    "full stack developer", "data scientist", "data analyst", "machine learning engineer",
    "ai engineer", "web developer", "product manager", "intern"
}
EDUCATION_TERMS = {
    "b.tech", "bachelor of technology", "b.e.", "bachelor of engineering", "b.sc",
    "bachelor of science", "m.tech", "master of technology", "m.sc", "master of science",
    "mba", "ph.d", "high school", "senior secondary", "secondary school"
}
DURATION_RE = re.compile(
    r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|"
    r"aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?|\d{4})"
    r"\s*(?:[-–—]|to)\s*(?:present|current|"
    r"jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|"
    r"aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?|\d{4})\b",
    re.IGNORECASE
)


def extract_text(filename: str, content: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        with fitz.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)
    elif lower.endswith(".docx"):
        doc = Document(BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raise ValueError("Only PDF and DOCX resumes are supported.")
    text = re.sub(r"[ \t]+", " ", text)
    if not text.strip():
        raise ValueError("No selectable text was found. This may be a scanned resume; OCR is not enabled yet.")
    return text[:200_000]


def _matches(text: str, terms: set[str]) -> list[str]:
    lower = text.lower()
    found = []
    for term in sorted(terms, key=len, reverse=True):
        pattern = r"(?<![a-z0-9])" + re.escape(term.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(term)
    return sorted(set(found))


def parse_resume(text: str) -> dict:
    # spaCy is used for sentence/entity processing when the model is installed.
    # Phrase matching is also used because general-purpose NER models do not reliably
    # recognize technical skills; this makes the extraction deterministic and editable.
    try:
        nlp = spacy.load("en_core_web_sm")
        doc = nlp(text[:100_000])
        entities = [(ent.text.strip(), ent.label_) for ent in doc.ents]
    except (OSError, ValueError):
        entities = []

    skills = _matches(text, SKILL_TERMS)
    titles = _matches(text, TITLE_TERMS)
    institutions = sorted({
        value for value, label in entities
        if label == "ORG" and any(word in value.lower() for word in (
            "university", "college", "institute", "school", "academy"
        ))
    })
    # Education labels are useful signals, but do not pretend every school is an NER hit.
    institutions = sorted(set(institutions + _matches(text, EDUCATION_TERMS)))
    durations = sorted(set(match.group(0).strip() for match in DURATION_RE.finditer(text)))
    return {
        "skills": skills,
        "job_titles": titles,
        "institutions": institutions,
        "duration_markers": durations,
    }
